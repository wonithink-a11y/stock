#!/usr/bin/env python3
"""섹터별 연도·월 성과 순위 + 달력월 계절성 + 돈의 쏠림(거래대금 비중·외국인/기관 순매수) 엑셀·차트, 그리고 쏠림 신호 검증.
설계는 실행 전 고정(2026-09-30). 자문 아님.

    python research/strategy-lab/sector_flow_calendar_excel.py --selftest
    python research/strategy-lab/sector_flow_calendar_excel.py
    -> reports/2026-09-sector-flows/sector-flows-yearly.xlsx · findings/sector-flow-signal-results-2026-09.json

[쏠림 신호 검증 — 사전 고정]
- 한국 20그룹(`config/sectorGroups.json`), 월초 신호일 t 의 값으로 그 달 수익(fwd1m, 그룹 종목 평균)을 예측하는가. 그룹당 5종목 미만인 달·그룹 제외.
- 신호 2개(가설 2개): (1) DTV3 = 그룹 거래대금 비중(20일 평균 거래대금 합 / 전체)의 3개월 전 대비 변화(%p)  (2) FI20 = (외국인+기관) 20일 순매수 / 20일 거래대금.
  (거래대금은 매수+매도 합이라 순유입이 아니다 — 그래서 (2)를 함께 본다.)
- 지표: 월별 순위상관(Spearman) IC(그룹 12개 이상인 달) 의 Newey-West(lag 3) t, 구간 TRAIN/VALID/TEST = 패널의 period 열. 난수 바닥선 = 달 안에서 fwd 를 섞어 전 절차 500회,
  2신호 x 3구간 |t| 최댓값의 p95. 판정: 바닥선을 넘고 세 구간 부호가 같으면 "신호 있음", 아니면 "없음/약함". 결과를 보고 창(3개월)·신호를 바꿔 재시험하지 않는다.
- 기록 전용(판정 아님): 같은 달 수급(다음 신호일의 20일 값)과 그 달 수익의 동시 상관 — "돈이 몰린 달에 올랐나".

[돈의 생애주기 — 사전 고정] 사용자 질문: 돈이 몰렸다가 어떻게 빠져나가는지, 들어가면 언제 나오는 게 유리한가.
- 수급 F = FI20(외국인+기관 20일 순매수/거래대금). 진입 사건 = 월초 신호일 t 에 그룹의 F 순위가 새로 상위 3 에 든 것(t-1 에는 상위 3 밖, 그룹 12개 이상인 달).
- 이벤트 시간 프로파일(기술통계, k=-3..+6): 평균 F, 평균 F 순위, 계속 상위 3 인 비율(귀무 3/그룹수), 진입 전·후 누적 초과수익(그룹 수익 - 그 달 전체 평균).
- 퇴출 규칙 8개 고정: 보유 1·2·3·6·12개월, F<0 이 되는 첫 달, F 순위가 10 밖으로 밀리는 첫 달, 거래대금 비중 3개월 변화(DTV3)<0 이 되는 첫 달. 최대 12개월. 결정은 t+h 신호일 값으로(다음 달 수익부터 반영, 룩어헤드 없음).
- 지표: 사건별 월평균 초과수익(복리 초과 / 보유월수) 평균 · 승률 · 평균 보유월. 12개월 블록 부트스트랩 2,000회. 난수 진입 플라시보(같은 수의 무작위 그룹-월, 300회)의 p95 를 바닥선으로.
  진입 t+12 자료가 있는 사건만(전 규칙 같은 표본). 판정: (a) 진입 자체가 우위 = 보유 3개월 규칙의 평균이 플라시보 p95 초과이고 부트스트랩 하한 > 0. (b) 수급 기반 퇴출이 낫다 = 그 규칙 - 보유 3개월의 짝지은 차이 하한 > 0.
  둘 다 아니면 "패턴은 기술할 수 있으나 수치로 쓸 우위는 없다". 규칙·창을 결과 보고 바꾸지 않는다.
- 한계: 사건 수 수백 건(서로 겹침), 수급은 2026-06 까지, 미국은 수급 자료 없음.
- 한계: 126개월, 그룹 = 현재 업종분류·생존 종목, 수급 자료 2026-06 까지. 소부장은 별도 그룹이 아니다('반도체' 그룹은 KSIC 코드 1개).
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

LAB = Path(__file__).resolve().parent
sys.path.insert(0, str(LAB))
from pension_sector_bond_test import mret  # noqa: E402
from sector_step0 import load_rollup, nw_tstat  # noqa: E402
from sector_leadership_macro_excel import US_NAMES, leaders, to_period  # noqa: E402

PANEL = LAB / "data" / "factor-panel" / "kr-monthly-v1.parquet"
OUT = LAB / "reports" / "2026-09-sector-flows"
RES_JSON = LAB / "findings" / "sector-flow-signal-results-2026-09.json"
N_FLOOR = 500
PER = ["TRAIN", "VALID", "TEST"]


def group_table():
    p = pd.read_parquet(PANEL, columns=["ticker", "date", "sector", "period", "dv20", "foreign_nb20_ratio", "inst_nb20_ratio",
                                        "indiv_nb20_ratio", "fwd1m"])
    p["date"] = pd.to_datetime(p["date"])
    p["group"] = p["sector"].map(load_rollup())
    p = p[p["group"].notna()].copy()
    p["amt20"] = p["dv20"] * 20.0
    for k, c in (("F", "foreign_nb20_ratio"), ("I", "inst_nb20_ratio"), ("P", "indiv_nb20_ratio")):
        p["net" + k] = p[c] * p["amt20"]
    agg = p.groupby(["date", "group"]).agg(n=("ticker", "size"), period=("period", "first"), amt20=("amt20", "sum"), netF=("netF", "sum"),
                                           netI=("netI", "sum"), netP=("netP", "sum"), fwd=("fwd1m", "mean")).reset_index()
    agg = agg[agg["n"] >= 5].copy()
    agg["share"] = agg["amt20"] / agg.groupby("date")["amt20"].transform("sum")
    agg["FI20"] = (agg["netF"] + agg["netI"]) / agg["amt20"]
    return agg


def add_dtv3(agg: pd.DataFrame):
    sh = agg.pivot(index="date", columns="group", values="share").sort_index()
    d3 = sh - sh.shift(3)
    agg = agg.merge(d3.stack().rename("DTV3").reset_index(), on=["date", "group"], how="left")
    return agg


def ic_series(d: pd.DataFrame, sig: str, shuffle_rng=None):
    ics = {}
    for dt, g in d.groupby("date"):
        g = g.dropna(subset=[sig, "fwd"])
        if len(g) < 12:
            continue
        y = g["fwd"].values
        if shuffle_rng is not None:
            y = shuffle_rng.permutation(y)
        ics[dt] = pd.Series(g[sig].values).corr(pd.Series(y), method="spearman")
    return pd.Series(ics)


def per_period(ic: pd.Series, per_map: dict):
    out = {}
    for p in PER:
        x = ic[[d for d in ic.index if per_map.get(d) == p]]
        out[p] = dict(n=int(len(x)), mean_ic=float(x.mean()), t=float(nw_tstat(x.values)))
    return out


def signal_test(agg: pd.DataFrame):
    per_map = agg.groupby("date")["period"].first().to_dict()
    sigs = ["DTV3", "FI20"]
    res = {}
    for s in sigs:
        res[s] = per_period(ic_series(agg, s), per_map)
    rng = np.random.default_rng(11)
    mx = []
    for _ in range(N_FLOOR):
        ts = []
        for s in sigs:
            pp = per_period(ic_series(agg, s, rng), per_map)
            ts.extend(abs(v["t"]) for v in pp.values() if np.isfinite(v["t"]))
        mx.append(max(ts))
    floor = float(np.quantile(mx, 0.95))
    verdict = {}
    for s in sigs:
        t = [res[s][p]["t"] for p in PER]
        ok = all(abs(x) > floor for x in t) and len({np.sign(x) for x in t}) == 1
        verdict[s] = "신호 있음" if ok else "없음/약함"
    # 기록 전용: 동시 상관(같은 달 수급 vs 그 달 수익)
    dates = sorted(agg["date"].unique())
    nxt = {dates[i]: dates[i + 1] for i in range(len(dates) - 1)}
    fi_next = agg.set_index(["date", "group"])["FI20"]
    c = agg.copy()
    c["FI_next"] = [fi_next.get((nxt.get(d), g), np.nan) for d, g in zip(c["date"], c["group"])]
    contemp = per_period(ic_series(c, "FI_next"), per_map)
    return dict(signals=res, floor_t=floor, verdict=verdict, contemporaneous_FI=contemp, n_months=int(agg["date"].nunique()))



# ---------------------------------------------------------------- 돈의 생애주기
RULES = ["보유1", "보유2", "보유3", "보유6", "보유12", "F<0", "F순위>10", "DTV3<0"]
MAXH = 12


def mats(agg: pd.DataFrame):
    dates = sorted(agg["date"].unique())
    pv = lambda c: agg.pivot(index="date", columns="group", values=c).reindex(dates)
    F, S, D, X = pv("FI20"), pv("share"), pv("DTV3"), pv("fwd")
    cols = list(F.columns)
    rk = F.rank(axis=1, ascending=False, method="first")
    rk[F.notna().sum(axis=1) < 12] = np.nan
    return dates, cols, F.values, S.values, D.values, X.values, rk.values


def hold_months(rule, t, g, F, D, rk, T):
    if rule.startswith("보유"):
        return int(rule[2:])
    for h in range(1, MAXH + 1):
        if t + h >= T:
            return h
        if rule == "F<0" and not (F[t + h, g] >= 0):
            return h
        if rule == "F순위>10" and not (rk[t + h, g] <= 10):
            return h
        if rule == "DTV3<0" and not (D[t + h, g] >= 0):
            return h
    return MAXH


def episode_excess(X, t, g, h):
    seg = X[t:t + h]
    if np.isnan(seg[:, g]).any():
        return np.nan
    bench = np.nanmean(seg, axis=1)
    return (np.prod(1 + seg[:, g]) - np.prod(1 + bench)) / h


def run_rules(entries, X, F, D, rk, T):
    out = {r: [] for r in RULES}
    hh = {r: [] for r in RULES}
    for t, g in entries:
        for r in RULES:
            h = hold_months(r, t, g, F, D, rk, T)
            out[r].append(episode_excess(X, t, g, h))
            hh[r].append(h)
    return {r: np.array(v, float) for r, v in out.items()}, {r: np.array(v) for r, v in hh.items()}


def lifecycle(agg: pd.DataFrame):
    dates, cols, F, S, D, X, rk = mats(agg)
    T, N = F.shape
    ent = []
    for t in range(1, T - MAXH - 1):
        for g in range(N):
            if rk[t, g] <= 3 and not (rk[t - 1, g] <= 3):
                ent.append((t, g))
    valid = [(t, g) for t in range(1, T - MAXH - 1) for g in range(N) if not np.isnan(rk[t, g])]
    prof = {}
    xs = X - np.nanmean(X, axis=1, keepdims=True)
    for k in range(-3, 7):
        f, r, top, cx, sd = [], [], [], [], []
        for t, g in ent:
            if 0 <= t + k < T:
                f.append(F[t + k, g])
                r.append(rk[t + k, g])
                top.append(float(rk[t + k, g] <= 3))
                sd.append(S[t + k, g] - S[t - 1, g])
            if k >= 0 and t + k < T:
                seg = xs[t:t + k + 1, g]
                cx.append(np.prod(1 + seg) - 1 if not np.isnan(seg).any() else np.nan)
            if k < 0 and t + k >= 0:
                seg = xs[t + k:t, g]
                cx.append(np.prod(1 + seg) - 1 if not np.isnan(seg).any() else np.nan)
        prof[k] = dict(F=float(np.nanmean(f)), rank=float(np.nanmean(r)), still_top3=float(np.nanmean(top)),
                       d_share_pp=float(np.nanmean(sd) * 100), cum_excess=float(np.nanmean(cx)), n=len(f))
    ex, hh = run_rules(ent, X, F, D, rk, T)
    rng = np.random.default_rng(21)
    ent_t = np.array([t for t, g in ent])
    B = 2000
    res = {}
    base3 = ex["보유3"]
    boots = {r: [] for r in RULES}
    diffs = {r: [] for r in RULES}
    nb = int(np.ceil(T / 12))
    for _ in range(B):
        st = rng.integers(0, T, nb)
        months = np.concatenate([(s0 + np.arange(12)) % T for s0 in st])[:T]
        cnt = np.bincount(months, minlength=T)
        w = cnt[ent_t].astype(float)
        if w.sum() == 0:
            continue
        for r in RULES:
            ok = ~np.isnan(ex[r]) & ~np.isnan(base3)
            ww = w * ok
            if ww.sum() == 0:
                continue
            boots[r].append(np.nansum(ex[r] * ww) / ww.sum())
            diffs[r].append(np.nansum((ex[r] - base3) * ww) / ww.sum())
    P = 300
    plc = {r: [] for r in RULES}
    for _ in range(P):
        idx = rng.choice(len(valid), size=len(ent), replace=False)
        e2, _h = run_rules([valid[i] for i in idx], X, F, D, rk, T)
        for r in RULES:
            plc[r].append(np.nanmean(e2[r]))
    for r in RULES:
        v = ex[r][~np.isnan(ex[r])]
        res[r] = dict(n=int(len(v)), mean_bp_per_month=float(v.mean() * 1e4),
                      ci=[float(np.quantile(boots[r], .025) * 1e4), float(np.quantile(boots[r], .975) * 1e4)],
                      hit=float((v > 0).mean()), avg_hold=float(hh[r].mean()), placebo_p95_bp=float(np.quantile(plc[r], .95) * 1e4),
                      placebo_mean_bp=float(np.mean(plc[r]) * 1e4),
                      diff_vs_hold3_bp=(None if r == "보유3" else float(np.mean(diffs[r]) * 1e4)),
                      diff_ci=(None if r == "보유3" else [float(np.quantile(diffs[r], .025) * 1e4), float(np.quantile(diffs[r], .975) * 1e4)]))
    entry_edge = res["보유3"]["mean_bp_per_month"] > res["보유3"]["placebo_p95_bp"] and res["보유3"]["ci"][0] > 0
    better = [r for r in RULES if r.startswith(("F", "DTV")) and res[r]["diff_ci"] and res[r]["diff_ci"][0] > 0]
    dur = []
    for t, g in ent:
        h = 0
        while t + h < T and rk[t + h, g] <= 3:
            h += 1
        dur.append(h)
    return dict(n_entries=len(ent), n_groups=int(N), profile={str(k): v for k, v in prof.items()}, rules=res, entry_edge=bool(entry_edge),
                better_flow_exits=better, top3_run_median=float(np.median(dur)), top3_run_mean=float(np.mean(dur)), chance_top3=3.0 / N)


# ---------------------------------------------------------------- 엑셀
def fmt_pct(x):
    return f"{x:+.0%}" if pd.notna(x) else ""


def year_block(R: pd.DataFrame, year: int, names: dict):
    ps = [p for p in R.index if p.year == year]
    M = R.loc[ps]
    yr = (1 + M).prod() - 1
    order = yr.sort_values(ascending=False).index
    rank_m = M.rank(axis=1, ascending=False, method="first")
    return ps, M, yr, order, rank_m


def write_year_sheet(wb, title, R, year, names, chart_prefix):
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    ps, M, yr, order, rank_m = year_block(R, year, names)
    ws = wb.create_sheet(title)
    hdr = PatternFill("solid", fgColor="1F3864")
    ws["A1"] = f"{title} — 섹터별 월별 수익률(색: -15% 빨강 ~ +15% 초록), 연수익·연순위"
    ws["A1"].font = Font(bold=True, size=13)
    heads = ["섹터"] + [f"{m}월" for m in range(1, 13)] + ["연수익", "연순위"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=3, column=j, value=h)
        c.fill, c.font, c.alignment = hdr, Font(color="FFFFFF", bold=True), Alignment(horizontal="center")
    for i, g in enumerate(order):
        r = 4 + i
        ws.cell(row=r, column=1, value=names.get(g, g))
        for p in ps:
            c = ws.cell(row=r, column=1 + p.month, value=float(M.loc[p, g]))
            c.number_format = "0.0%"
        c = ws.cell(row=r, column=14, value=float(yr[g]))
        c.number_format = "0.0%"
        c.font = Font(bold=True)
        ws.cell(row=r, column=15, value=i + 1)
    n = len(order)
    ws.conditional_formatting.add(f"B4:N{3 + n}", ColorScaleRule(start_type="num", start_value=-0.15, start_color="F8696B", mid_type="num",
                                  mid_value=0, mid_color="FFFFFF", end_type="num", end_value=0.15, end_color="63BE7B"))
    ws.column_dimensions["A"].width = 20
    for j in range(2, 16):
        ws.column_dimensions[chr(64 + j)].width = 8.5
    # 월별 1·2·3위 표
    r0 = 5 + n
    ws.cell(row=r0, column=1, value="월별 1·2·3위 섹터 (그 달 수익 기준)").font = Font(bold=True, size=12)
    for j, h in enumerate(["월", "1위", "수익", "2위", "수익", "3위", "수익", "꼴찌", "수익"], 1):
        c = ws.cell(row=r0 + 1, column=j, value=h)
        c.fill, c.font, c.alignment = hdr, Font(color="FFFFFF", bold=True), Alignment(horizontal="center")
    for k, p in enumerate(ps):
        row = M.loc[p].sort_values(ascending=False)
        vals = [f"{p.month}월", names.get(row.index[0], row.index[0]), row.iloc[0], names.get(row.index[1], row.index[1]), row.iloc[1],
                names.get(row.index[2], row.index[2]), row.iloc[2], names.get(row.index[-1], row.index[-1]), row.iloc[-1]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r0 + 2 + k, column=j, value=v)
            if j in (3, 5, 7, 9):
                c.number_format = "0.0%"
    ws.column_dimensions["D"].width = 14
    # 차트용 누적 표(연수익 상위 5)
    top5 = list(order[:5])
    c0 = r0 + 4 + len(ps)
    ws.cell(row=c0, column=1, value="누적수익(연수익 상위 5, 1월 시작=100)").font = Font(bold=True)
    ws.cell(row=c0 + 1, column=1, value="월")
    for j, g in enumerate(top5, 2):
        ws.cell(row=c0 + 1, column=j, value=names.get(g, g))
    for k, p in enumerate(ps):
        ws.cell(row=c0 + 2 + k, column=1, value=f"{p.month}월")
        for j, g in enumerate(top5, 2):
            v = 100 * float((1 + M.loc[ps[:k + 1], g]).prod())
            ws.cell(row=c0 + 2 + k, column=j, value=round(v, 1))
    bar = BarChart()
    bar.type = "bar"
    bar.title = f"{year}년 섹터별 연수익"
    bar.add_data(Reference(ws, min_col=14, min_row=3, max_row=3 + n), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + n))
    bar.legend = None
    bar.height, bar.width = 12, 16
    bar.y_axis.number_format = "0%"
    ws.add_chart(bar, "Q3")
    ln = LineChart()
    ln.title = f"{year}년 연수익 상위 5 섹터 누적(=100)"
    ln.add_data(Reference(ws, min_col=2, max_col=6, min_row=c0 + 1, max_row=c0 + 1 + len(ps)), titles_from_data=True)
    ln.set_categories(Reference(ws, min_col=1, min_row=c0 + 2, max_row=c0 + 1 + len(ps)))
    ln.height, ln.width = 9, 16
    ws.add_chart(ln, "Q29")


def heat_sheet(wb, title, df, fmt, lo, hi, diverging=True, note=None, first_width=11):
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(title)
    r = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(bold=True)
        r = 3
    hdr = PatternFill("solid", fgColor="1F3864")
    cols = [df.index.name or "구분"] + list(df.columns)
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=r, column=j, value=h)
        c.fill, c.font, c.alignment = hdr, Font(color="FFFFFF", bold=True), Alignment(horizontal="center", wrap_text=True)
    for i, (idx, row) in enumerate(df.iterrows(), r + 1):
        ws.cell(row=i, column=1, value=str(idx))
        for j, v in enumerate(row.values, 2):
            if pd.notna(v):
                c = ws.cell(row=i, column=j, value=float(v) if isinstance(v, (int, float, np.floating, np.integer)) else v)
                if fmt:
                    c.number_format = fmt
    ws.freeze_panes = ws.cell(row=r + 1, column=2)
    last = r + len(df)
    from openpyxl.utils import get_column_letter
    rng = f"B{r + 1}:{get_column_letter(len(cols))}{last}"
    if diverging:
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=lo, start_color="F8696B", mid_type="num",
                                      mid_value=(lo + hi) / 2, mid_color="FFFFFF", end_type="num", end_value=hi, end_color="63BE7B"))
    else:
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=lo, start_color="63BE7B", mid_type="num",
                                      mid_value=(lo + hi) / 2, mid_color="FFFFFF", end_type="num", end_value=hi, end_color="F8696B"))
    ws.column_dimensions["A"].width = first_width
    for j in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 11
    ws.row_dimensions[r].height = 32
    return ws


def calendar_tables(R: pd.DataFrame, names: dict):
    rk = R.rank(axis=1, ascending=False, method="first")
    mon = R.index.month
    avg = R.groupby(mon).mean().T
    cnt1 = (rk == 1).groupby(mon).sum().T
    cnt3 = (rk <= 3).groupby(mon).sum().T
    avgrk = rk.groupby(mon).mean().T
    for t in (avg, cnt1, cnt3, avgrk):
        t.index = [names.get(i, i) for i in t.index]
        t.columns = [f"{m}월" for m in t.columns]
    yrs = sorted(set(R.index.year))
    n1 = pd.DataFrame(index=yrs, columns=[f"{m}월" for m in range(1, 13)], dtype=object)
    n2 = n1.copy()
    for p in R.index:
        s = R.loc[p].sort_values(ascending=False)
        n1.loc[p.year, f"{p.month}월"] = f"{names.get(s.index[0], s.index[0])} {s.iloc[0]:+.0%}"
        n2.loc[p.year, f"{p.month}월"] = f"{names.get(s.index[1], s.index[1])} {s.iloc[1]:+.0%}"
    return avg, cnt1, cnt3, avgrk, n1, n2


def text_sheet(wb, title, dfs: list):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(title)
    hdr = PatternFill("solid", fgColor="1F3864")
    r = 1
    for name, df in dfs:
        ws.cell(row=r, column=1, value=name).font = Font(bold=True, size=12)
        r += 1
        for j, h in enumerate(["연도"] + list(df.columns), 1):
            c = ws.cell(row=r, column=j, value=h)
            c.fill, c.font, c.alignment = hdr, Font(color="FFFFFF", bold=True), Alignment(horizontal="center")
        for idx, row in df.iterrows():
            r += 1
            ws.cell(row=r, column=1, value=int(idx))
            for j, v in enumerate(row.values, 2):
                ws.cell(row=r, column=j, value=v if isinstance(v, str) else None)
        r += 3
    ws.column_dimensions["A"].width = 8
    for j in range(2, 14):
        ws.column_dimensions[chr(64 + j)].width = 20
    return ws


def calendar_sheet(wb, title, R, names):
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    avg, cnt1, cnt3, avgrk, n1, n2 = calendar_tables(R, names)
    ws = wb.create_sheet(title)
    hdr = PatternFill("solid", fgColor="1F3864")
    r = 1
    specs = [("① 달력월별 평균 월수익 (여러 해 평균)", avg, "0.0%", -0.05, 0.05, True),
             ("② 그 달력월에 1위였던 횟수(해)", cnt1, "0", 0, max(3, int(cnt1.values.max())), True),
             ("③ 그 달력월에 상위 3위 안이었던 횟수(해)", cnt3, "0", 0, max(3, int(cnt3.values.max())), True),
             ("④ 달력월별 평균 순위 (1이 최고, 낮을수록 좋음)", avgrk, "0.0", 1, R.shape[1], False)]
    for name, df, fmt, lo, hi, div in specs:
        ws.cell(row=r, column=1, value=name).font = Font(bold=True, size=12)
        r += 1
        for j, h in enumerate(["섹터"] + list(df.columns), 1):
            c = ws.cell(row=r, column=j, value=h)
            c.fill, c.font, c.alignment = hdr, Font(color="FFFFFF", bold=True), Alignment(horizontal="center")
        top = r + 1
        for idx, row in df.iterrows():
            r += 1
            ws.cell(row=r, column=1, value=idx)
            for j, v in enumerate(row.values, 2):
                c = ws.cell(row=r, column=j, value=float(v))
                c.number_format = fmt
        rng = f"B{top}:{get_column_letter(1 + df.shape[1])}{r}"
        if div:
            ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=lo, start_color="F8696B" if lo < 0 else "FFFFFF",
                                          mid_type="num", mid_value=(lo + hi) / 2 if lo < 0 else (lo + hi) / 2, mid_color="FFFFFF",
                                          end_type="num", end_value=hi, end_color="63BE7B"))
        else:
            ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=lo, start_color="63BE7B", end_type="num",
                                          end_value=hi, end_color="FFFFFF"))
        r += 3
    ws.column_dimensions["A"].width = 20
    for j in range(2, 14):
        ws.column_dimensions[get_column_letter(j)].width = 9
    t = text_sheet(wb, title + "_1위2위표", [("연도 x 달력월: 그 달 1위 섹터(수익)", n1), ("연도 x 달력월: 그 달 2위 섹터(수익)", n2)])
    return ws, t


def flow_sheets(wb, agg: pd.DataFrame):
    from openpyxl.chart import AreaChart, Reference
    sh = agg.pivot(index="date", columns="group", values="share").sort_index()
    sh.index = [str(pd.Period(d, "M") - 1) for d in sh.index]     # 흐름이 측정된 달(신호일 직전 20세션)
    sh.index.name = "유입월"
    ws = heat_sheet(wb, "한국_거래대금비중", sh, "0.0%", 0.0, 0.25, diverging=False,
                    note="그룹별 20일 거래대금 비중(전체=100%). 유입월 = 신호일 직전 20세션이 속한 달. 거래대금은 매수+매도 합이라 순유입이 아니다.")
    # 차트용: 최근 24개월 평균 비중 상위 8 + 기타
    top = sh.tail(24).mean().sort_values(ascending=False).index[:8]
    cd = sh[top].copy()
    cd["기타"] = 1 - cd.sum(axis=1)
    c0 = 4 + len(sh) + 3
    ws.cell(row=c0, column=1, value="차트용(최근 24개월 평균 비중 상위 8 + 기타)")
    ws.cell(row=c0 + 1, column=1, value="유입월")
    for j, g in enumerate(cd.columns, 2):
        ws.cell(row=c0 + 1, column=j, value=g)
    for i, (idx, row) in enumerate(cd.iterrows()):
        ws.cell(row=c0 + 2 + i, column=1, value=idx)
        for j, v in enumerate(row.values, 2):
            ws.cell(row=c0 + 2 + i, column=j, value=float(v)).number_format = "0.0%"
    ch = AreaChart()
    ch.grouping = "percentStacked"
    ch.title = "섹터별 거래대금 비중 추이(상위 8 + 기타)"
    ch.add_data(Reference(ws, min_col=2, max_col=1 + cd.shape[1], min_row=c0 + 1, max_row=c0 + 1 + len(cd)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=c0 + 2, max_row=c0 + 1 + len(cd)))
    ch.height, ch.width = 12, 26
    ws.add_chart(ch, f"B{c0 + len(cd) + 4}")
    for key, lab in (("netF", "외국인"), ("netI", "기관"), ("netP", "개인")):
        t = (agg.pivot(index="date", columns="group", values=key).sort_index() / 1e8)
        t.index = [str(pd.Period(d, "M") - 1) for d in t.index]
        t.index.name = "유입월"
        lim = float(np.nanpercentile(np.abs(t.values), 90))
        heat_sheet(wb, f"한국_{lab}순매수(억)", t, "#,##0", -lim, lim, note=f"그룹별 {lab} 20일 순매수 추정액(억원) = 20일 순매수비율 x 20일 거래대금 합. 초록=순매수, 빨강=순매도.")
    fi = agg.pivot(index="date", columns="group", values="FI20").sort_index()
    fi.index = [str(pd.Period(d, "M") - 1) for d in fi.index]
    fi.index.name = "유입월"
    heat_sheet(wb, "한국_외국인+기관(%거래대금)", fi, "0.0%", -0.10, 0.10, note="(외국인+기관) 20일 순매수 / 20일 거래대금. 초록=순매수 우위.")


def life_sheet(wb, life):
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    ws = wb.create_sheet("돈의생애주기")
    ws["A1"] = f"수급 생애주기 — 외국인+기관 순매수 순위가 새로 상위 3에 든 사건 {life['n_entries']}건 (한국 {life['n_groups']}그룹, 월초 신호일 기준). k=0 이 진입 신호일, 수익은 k=0 달부터."
    ws["A1"].font = Font(bold=True)
    hdr = PatternFill("solid", fgColor="1F3864")
    heads = ["k(월)", "평균 FI20(순매수/거래대금)", "평균 F순위", "계속 상위3 비율", "거래대금비중 변화(진입 전월 대비 %p)", "누적 초과수익(진입 전: 직전 |k|개월, 진입 후: k+1개월)", "사건수"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=3, column=j, value=h)
        c.fill, c.font = hdr, Font(color="FFFFFF", bold=True)
    prof = sorted(life["profile"].items(), key=lambda kv: int(kv[0]))
    for i, (k, v) in enumerate(prof, 4):
        for j, x in enumerate([int(k), v["F"], v["rank"], v["still_top3"], v["d_share_pp"], v["cum_excess"], v["n"]], 1):
            c = ws.cell(row=i, column=j, value=x)
            if j in (2, 4, 6):
                c.number_format = "0.0%"
            elif j in (3, 5):
                c.number_format = "0.0"
    r0 = 4 + len(prof) + 2
    ws.cell(row=r0, column=1, value=f"상위3 유지 기간: 평균 {life['top3_run_mean']:.1f}개월 · 중앙 {life['top3_run_median']:.0f}개월 (무작위 상위3 확률 {life['chance_top3']:.0%})").font = Font(bold=True)
    r0 += 2
    heads = ["퇴출 규칙", "사건수", "월평균 초과수익(bp)", "95% 구간 하한", "95% 구간 상한", "승률", "평균 보유(개월)", "난수 진입 평균(bp)", "난수 진입 p95(bp)", "보유3 대비 차이(bp)", "차이 하한", "차이 상한"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.fill, c.font = hdr, Font(color="FFFFFF", bold=True)
    for i, r in enumerate(RULES, r0 + 1):
        v = life["rules"][r]
        vals = [r, v["n"], v["mean_bp_per_month"], v["ci"][0], v["ci"][1], v["hit"], v["avg_hold"], v["placebo_mean_bp"], v["placebo_p95_bp"],
                v["diff_vs_hold3_bp"], (v["diff_ci"][0] if v["diff_ci"] else None), (v["diff_ci"][1] if v["diff_ci"] else None)]
        for j, x in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=x)
            if j == 6:
                c.number_format = "0%"
            elif j in (3, 4, 5, 8, 9, 10, 11, 12) and x is not None:
                c.number_format = "0.0"
    end = r0 + len(RULES) + 2
    ws.cell(row=end, column=1, value=f"판정: 진입 자체 우위 = {'있음' if life['entry_edge'] else '없음'} · 수급 기반 퇴출이 보유3 보다 나음 = {life['better_flow_exits'] or '없음'}").font = Font(bold=True, size=12)
    ln = LineChart()
    ln.title = "진입 전후 평균 누적 초과수익"
    ln.add_data(Reference(ws, min_col=6, min_row=3, max_row=3 + len(prof)), titles_from_data=True)
    ln.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + len(prof)))
    ln.height, ln.width = 9, 18
    ln.y_axis.number_format = "0.0%"
    ws.add_chart(ln, "N3")
    ln2 = LineChart()
    ln2.title = "진입 전후 평균 FI20 (순매수/거래대금)"
    ln2.add_data(Reference(ws, min_col=2, min_row=3, max_row=3 + len(prof)), titles_from_data=True)
    ln2.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + len(prof)))
    ln2.height, ln2.width = 9, 18
    ln2.y_axis.number_format = "0.0%"
    ws.add_chart(ln2, "N22")
    ws.column_dimensions["A"].width = 14
    for j in range(2, 13):
        ws.column_dimensions[chr(64 + j)].width = 16


def current_snapshot(wb):
    from openpyxl.styles import Font, PatternFill
    d = json.load(open(LAB.parents[1] / "docs" / "data" / "sector-strength.json", encoding="utf-8"))
    rows = []
    for g in d["groups"]:
        rows.append({"섹터": g["group"], "종목수": g["n"], "1개월": g["ret"]["1m"], "3개월": g["ret"]["3m"], "6개월": g["ret"]["6m"],
                     "3개월 상대강도(RS)": g["rs"]["3m"], "거래대금가중 3개월": g.get("tvRet", {}).get("3m"), "4분면": g["quadrant"],
                     "20일선 위 비율": g["breadth20"]})
    df = pd.DataFrame(rows).sort_values("3개월", ascending=False)
    ws = wb.create_sheet("현재(대시보드)")
    ws["A1"] = f"현재 섹터 강도 — 기준일 {d['asOf']} (docs/data/sector-strength.json, 종목 중앙값 수익). 순매수·거래대금 비중은 2026-06 까지만 있어 여기엔 없다."
    ws["A1"].font = Font(bold=True)
    for j, h in enumerate(df.columns, 1):
        c = ws.cell(row=3, column=j, value=h)
        c.fill, c.font = PatternFill("solid", fgColor="1F3864"), Font(color="FFFFFF", bold=True)
    for i, row in enumerate(df.itertuples(index=False), 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            if j in (3, 4, 5, 6, 7, 9):
                c.number_format = "0.0%"
    ws.column_dimensions["A"].width = 20
    return df


def main():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    agg = add_dtv3(group_table())
    test = signal_test(agg)
    life = lifecycle(agg)
    test["lifecycle"] = life
    RES_JSON.write_text(json.dumps(test, ensure_ascii=False, indent=1), encoding="utf-8")
    Rkr = agg.pivot(index="date", columns="group", values="fwd").sort_index()
    Rkr = Rkr.dropna(axis=1, thresh=int(len(Rkr) * 0.9)).dropna(how="any")
    Rkr = to_period(Rkr)
    Rus = to_period(pd.concat({s: mret(s) for s in US_NAMES}, axis=1).dropna()["1999-01":])
    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    notes = ["섹터별 연도·월 성과 순위, 달력월 계절성, 돈의 쏠림(거래대금 비중·외국인/기관 순매수). 자문 아님 — 과거 기록.",
             "한국 20그룹: 종목 평균 월수익, 2016-02~2026-07, 현재 업종분류·생존 종목 기준. 미국: SPDR 9섹터 총수익 1999-01~2026-09.",
             "시트: 한국_연도별섹터 / 한국_월수익 / 한국_월순위 / 한국_달력월(+_1위2위표) / 한국_거래대금비중(차트 포함) / 한국_외국인·기관·개인순매수 / 한국_외국인+기관(%거래대금) / KR2016~KR2026(연도별 월표·차트) / 미국_* / 쏠림검증 / 현재(대시보드).",
             "월 정의: 월수익은 그 달 첫 세션~다음 달 첫 세션. 수급 '유입월'은 신호일 직전 20세션이 속한 달(수급 자료는 2026-06 까지).",
             "★ 소부장(소재·부품·장비)은 별도 그룹이 아니다. '반도체' 그룹은 KSIC 반도체 코드 하나(1종 업종)이고 소부장 종목은 기계·장비, 전자부품·디스플레이, 화학·소재에 흩어져 있다.",
             "★ 거래대금은 매수+매도 합이라 순유입이 아니다. 순매수(외국인·기관)를 함께 본다. 어느 쪽도 '다음 달 수익'을 예측하는지는 '쏠림검증' 시트(사전 고정 설계)를 본다.",
             "★ 한국 표본 후반(2025~26)에 국내 증시 급등이 몰려 있다."]
    for i, t in enumerate(notes, 1):
        ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 160
    def yearly(R, names):
        yrs = sorted(set(R.index.year))
        t = pd.DataFrame({y: (1 + R[[p.year == y for p in R.index]]).prod() - 1 for y in yrs})
        t.index = [names.get(i, i) for i in t.index]
        t.columns = [f"{y}" for y in t.columns]
        return t
    for tag, R, names in (("한국", Rkr, {}), ("미국", Rus, US_NAMES)):
        yt = yearly(R, names)
        heat_sheet(wb, f"{tag}_연도별섹터", yt.round(4), "0.0%", -0.3, 0.3, note="섹터별 연수익(복리, 그 해 자료가 있는 달만). 색: -30% 빨강 ~ +30% 초록.", first_width=20)
        yr_rank = yt.rank(ascending=False, method="first").astype(int)
        yr_rank.index.name = "섹터 순위"
        heat_sheet(wb, f"{tag}_연도별순위", yr_rank, "0", 1, len(yr_rank), diverging=False, note="연수익 순위(1=최고).", first_width=20)
        mr = R.copy()
        mr.index = mr.index.astype(str)
        mr.columns = [names.get(c, c) for c in mr.columns]
        mr.index.name = "연월"
        heat_sheet(wb, f"{tag}_월수익", mr, "0.0%", -0.15, 0.15, note="섹터별 월수익(색: -15% ~ +15%).", first_width=10)
        rk = R.rank(axis=1, ascending=False, method="first").astype(int)
        rk.index = rk.index.astype(str)
        rk.columns = [names.get(c, c) for c in rk.columns]
        rk.index.name = "연월"
        heat_sheet(wb, f"{tag}_월순위", rk, "0", 1, rk.shape[1], diverging=False, note="그 달 수익 순위(1=최고, 초록).", first_width=10)
        calendar_sheet(wb, f"{tag}_달력월", R, names)
    for y in sorted(set(Rkr.index.year)):
        write_year_sheet(wb, f"KR{y}", Rkr, y, {}, "KR")
    ws = wb.create_sheet("쏠림검증")
    lines = [("쏠림 신호가 다음 달 섹터 수익을 예측하는가 (사전 고정 설계, 20그룹 x 월별 순위상관 IC)", None)]
    ws.cell(row=1, column=1, value=lines[0][0])
    r = 3
    for j, h in enumerate(["신호", "구간", "개월수", "평균 IC", "NW t", "난수 바닥선 t(p95)", "판정"], 1):
        ws.cell(row=r, column=j, value=h)
    for s, lab in (("DTV3", "거래대금 비중 3개월 변화(%p)"), ("FI20", "외국인+기관 20일 순매수/거래대금")):
        for p in PER:
            r += 1
            v = test["signals"][s][p]
            for j, x in enumerate([lab, p, v["n"], round(v["mean_ic"], 4), round(v["t"], 2), round(test["floor_t"], 2), test["verdict"][s]], 1):
                ws.cell(row=r, column=j, value=x)
    r += 2
    ws.cell(row=r, column=1, value="기록 전용 — 같은 달 수급(외국인+기관)과 그 달 수익의 동시 상관")
    for p in PER:
        r += 1
        v = test["contemporaneous_FI"][p]
        ws.cell(row=r, column=1, value=p)
        ws.cell(row=r, column=3, value=v["n"])
        ws.cell(row=r, column=4, value=round(v["mean_ic"], 4))
        ws.cell(row=r, column=5, value=round(v["t"], 2))
    ws.column_dimensions["A"].width = 40
    flow_sheets(wb, agg)
    life_sheet(wb, life)
    snap = current_snapshot(wb)
    OUT.mkdir(parents=True, exist_ok=True)
    p = OUT / "sector-flows-yearly.xlsx"
    wb.save(p)
    print("wrote", p, len(wb.sheetnames), "sheets")
    print(json.dumps(test, ensure_ascii=False)[:1500])


def selftest():
    rng = np.random.default_rng(0)
    idx = pd.period_range("2020-01", periods=24, freq="M")
    R = pd.DataFrame(rng.normal(0, .05, (24, 6)), index=idx, columns=list("abcdef"))
    avg, c1, c3, ar, n1, n2 = calendar_tables(R, {})
    assert avg.shape == (6, 12) and (c1.sum().sum() == 24) and (c3.sum().sum() == 72) and n1.shape == (2, 12)
    ps, M, yr, order, rk = year_block(R, 2020, {})
    assert len(ps) == 12 and len(order) == 6
    # 신호 검증 도구: 완전 예측 신호면 IC=1, 무관하면 |IC| 작다
    d = pd.DataFrame({"date": np.repeat(pd.date_range("2020-01-01", periods=10, freq="MS"), 15), "group": list(range(15)) * 10})
    d["fwd"] = rng.normal(size=len(d))
    d["S1"] = d["fwd"]
    d["S2"] = rng.normal(size=len(d))
    assert ic_series(d, "S1").mean() > 0.99 and abs(ic_series(d, "S2").mean()) < 0.2
    # 생애주기: 보유N 은 N 개월, 무관 데이터에서 수익 0 이면 초과 0
    ag = []
    for di, dt in enumerate(pd.date_range("2016-01-01", periods=60, freq="MS")):
        for g in range(15):
            ag.append(dict(date=dt, group=g, n=6, period="TRAIN", amt20=1.0, netF=0.0, netI=0.0, netP=0.0, fwd=rng.normal(0, .05), share=1 / 15, FI20=rng.normal(0, .05)))
    ag = pd.DataFrame(ag)
    shp = ag.pivot(index="date", columns="group", values="share")
    ag = ag.merge((shp - shp.shift(3)).stack().rename("DTV3").reset_index(), on=["date", "group"], how="left")
    dts, cols, F, S, D, X, rk = mats(ag)
    assert hold_months("보유6", 3, 2, F, D, rk, 60) == 6 and 1 <= hold_months("F<0", 3, 2, F, D, rk, 60) <= 12
    assert abs(episode_excess(np.zeros((10, 3)), 2, 1, 3)) < 1e-12
    print("selftest ok")


if __name__ == "__main__":
    selftest() if "--selftest" in sys.argv else main()

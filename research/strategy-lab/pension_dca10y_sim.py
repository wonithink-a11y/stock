#!/usr/bin/env python3
"""연금저축 월 납입 10년 시뮬레이션 — 실행 전 고정 설계(2026-09-30). 자문 아님, 과거 데이터 기반 시나리오.

    python research/strategy-lab/pension_dca10y_sim.py --selftest
    python research/strategy-lab/pension_dca10y_sim.py
    -> findings/pension-dca10y-results-2026-09.json · reports/2026-09-pension-dca10y/pension-dca10y.xlsx

- 납입: 매월 초 50만원 x 120개월 = 원금 6,000만원(연 600만원 = 연금저축 세액공제 한도). 금액은 비례하므로 결과는 배수(최종액/원금)와 원(6,000만원 기준)으로 낸다.
- 배분 6개(앞선 시험과 동일, 결과 전 고정): A 주식100 · B 주식70/장기채30 · C 60/40(장기채) · D 60/40(중기채) · E 60/40(현금) · F 60/장기채30/금10.
  주식=SPY, 장기채=30년 par 채권 프록시(^TYX 정밀가격, TLT 상관 0.991), 중기채=10년 par 프록시(^TNX), 현금=^IRX, 금=GLD(2004-12~). USD 로컬 총수익, 매수 편도 5bp.
- 운용 방식: 납입금을 목표 비중보다 부족한 자산에 배분(팔지 않음). 리밸런스 매도는 하지 않는다(앞선 시험에서 납입형과 연 1회 매도 리밸런스는 구분되지 않음).
- 시험 4개(결과 전 고정):
  (1) 역사 롤링 10년 창: 1993-02~ 매달 시작하는 모든 창(F 는 2004-12~). 지표: 최종액/원금 최악·하위10%·중앙·상위10%·최고, 원금 미만 확률, 납입 가중 연수익(IRR) 중앙·최악, 납입 기간 중 평가액/누적 납입 최저.
  (2) 블록 부트스트랩: 같은 자료의 월수익 행(자산 동시)을 12개월 블록으로 5,000경로(A~E 는 1993-02~, F 는 2004-12~ 자료). 최종액/원금 분포. 블록 구조는 장기 추세·금리 체제를 보존하지 못한다는 한계.
  (3) 시작 시점별 실제 사례: 2000-01 · 2007-10 · 2010-01 · 2013-01 · 2016-10(=가장 최근 완결 10년, 2026-09 까지).
  (4) 30년물 금리 시나리오(미래 지향 민감도): 채권 다리만 금리 경로로 바꾼다 — 현재 5.59% 에서 120개월에 걸쳐 직선으로 3.5%·5.59%(불변)·7.5%·9.0% 로 이동, 채권 월수익=정밀가격, 주식은 (1)의 실제 창 수익 사용. 배분 C 만.
- 환 민감도(기록 전용): 2004-01~ 창에서 SPY·장기채·금을 원화 비헤지(x 환율)로 A·B·C·F.
- 한계: 하나의 역사 경로(강세장 편중, 채권은 1980년대 이후 금리 하락기), 환·세금·수수료·분배 세부 미반영, 부트스트랩은 독립 블록 가정.
"""
from __future__ import annotations

import json
import math
import sys
from pathlib import Path

import numpy as np
import pandas as pd

LAB = Path(__file__).resolve().parent
sys.path.insert(0, str(LAB))
from pension_sector_bond_test import mret, monthly_last, par_bond_tr  # noqa: E402

COST = 0.0005
MONTHS = 120
MONTHLY = 500_000
PRINCIPAL = MONTHLY * MONTHS
ALLOC = {"A 주식100": {"SPY": 1.0}, "B 70/30(장기채)": {"SPY": .7, "B30": .3}, "C 60/40(장기채)": {"SPY": .6, "B30": .4},
         "D 60/40(중기채)": {"SPY": .6, "B10": .4}, "E 60/40(현금)": {"SPY": .6, "CASH": .4}, "F 60/장기채30/금10": {"SPY": .6, "B30": .3, "GLD": .1}}
COHORTS = ["2000-01", "2007-10", "2010-01", "2013-01", "2016-10"]
RES = LAB / "findings" / "pension-dca10y-results-2026-09.json"
OUT = LAB / "reports" / "2026-09-pension-dca10y"


def path(Rv: np.ndarray, tw: np.ndarray, start: int):
    """월초 납입 1(부족 자산에만 매수), 월말 평가액 배열 반환."""
    h = np.zeros(len(tw))
    V = np.zeros(MONTHS)
    for m in range(MONTHS):
        want = tw * (h.sum() + 1)
        buy = np.maximum(want - h, 0)
        buy = buy / buy.sum()
        h = h + buy * (1 - COST)
        h = h * (1 + Rv[start + m])
        V[m] = h.sum()
    return V


def irr(final: float) -> float:
    """월초 납입 1 x 120, 월말 평가 final 을 만드는 월수익률 -> 연환산."""
    lo, hi = -0.05, 0.10
    for _ in range(80):
        mid = (lo + hi) / 2
        fv = sum((1 + mid) ** (MONTHS - t) for t in range(MONTHS))
        lo, hi = (mid, hi) if fv < final else (lo, mid)
    return (1 + mid) ** 12 - 1


def summarize(V_list):
    finals = np.array([V[-1] for V in V_list])
    ratio = finals / MONTHS
    under = np.array([np.min(V[5:] / np.arange(6, MONTHS + 1)) for V in V_list])
    irrs = np.array([irr(f) for f in finals])
    q = lambda a, p: float(np.quantile(a, p))
    return dict(n=len(V_list), worst=float(ratio.min()), p10=q(ratio, .1), median=q(ratio, .5), p90=q(ratio, .9), best=float(ratio.max()),
                loss_prob=float((ratio < 1).mean()), irr_median=q(irrs, .5), irr_worst=float(irrs.min()), irr_p10=q(irrs, .1), underwater_min=float(under.min()),
                underwater_p10=q(under, .1))


def assets():
    b30 = par_bond_tr(monthly_last("_TYX"), 30.0)
    b10 = par_bond_tr(monthly_last("_TNX"), 10.0)
    cash = monthly_last("_IRX") / 100 / 12
    P = pd.concat([mret("SPY"), b30, b10, cash, mret("GLD")], axis=1, keys=["SPY", "B30", "B10", "CASH", "GLD"])
    P = P.dropna(subset=["SPY", "B30", "B10", "CASH"])["1993-02":]
    P.index = P.index.to_period("M")
    return P


def run():
    P = assets()
    res = {"meta": dict(principal=PRINCIPAL, monthly=MONTHLY, months=int(len(P)), first=str(P.index[0]), last=str(P.index[-1]))}
    rng = np.random.default_rng(41)
    paths_last = {}
    hist, boot, coh = {}, {}, {}
    for name, tw_d in ALLOC.items():
        cols = list(tw_d)
        sub = P[cols].dropna()
        Rv = sub.values
        tw = np.array([tw_d[c] for c in cols])
        starts = list(range(0, len(sub) - MONTHS + 1))
        vl = [path(Rv, tw, s) for s in starts]
        hist[name] = summarize(vl)
        hist[name]["first_start"], hist[name]["last_start"] = str(sub.index[0]), str(sub.index[starts[-1]])
        paths_last[name] = (vl[-1] / (np.arange(1, MONTHS + 1))).tolist()
        paths_last[name + "_value"] = (vl[-1] * MONTHLY).tolist()
        # 부트스트랩
        T = len(sub)
        bl = []
        nb = math.ceil(MONTHS / 12)
        for _ in range(5000):
            st = rng.integers(0, T, nb)
            idx = np.concatenate([(s0 + np.arange(12)) % T for s0 in st])[:MONTHS]
            bl.append(path(Rv[idx], tw, 0))
        boot[name] = summarize(bl)
        # 코호트
        c = {}
        for cs in COHORTS:
            p0 = pd.Period(cs, "M")
            if p0 in sub.index:
                s = sub.index.get_loc(p0)
                if s + MONTHS <= len(sub):
                    V = path(Rv, tw, s)
                    c[cs] = dict(final_ratio=float(V[-1] / MONTHS), final_krw=float(V[-1] * MONTHLY), end=str(sub.index[s + MONTHS - 1]),
                                 underwater_min=float(np.min(V[5:] / np.arange(6, MONTHS + 1))))
        coh[name] = c
    res["hist"], res["boot"], res["cohort"] = hist, boot, coh
    # (4) 금리 시나리오 (C 배분)
    y0 = float(monthly_last("_TYX").iloc[-1])
    tw_c = np.array([.6, .4])
    sub = P[["SPY", "B30"]]
    scen = {}
    idx = pd.period_range("2030-01", periods=MONTHS + 1, freq="M").to_timestamp("M")
    for yT in (3.5, y0, 7.5, 9.0):
        ypath = pd.Series(np.linspace(y0, yT, MONTHS + 1), index=idx)
        bret = par_bond_tr(ypath, 30.0).values
        vl, va = [], []
        for s in range(0, len(sub) - MONTHS + 1):
            spy = sub["SPY"].values[s:s + MONTHS]
            Rm = np.column_stack([spy, bret])
            vl.append(path(Rm, tw_c, 0))
            va.append(path(spy.reshape(-1, 1), np.array([1.0]), 0))
        scen[f"{yT:.2f}"] = dict(C=summarize(vl), A=summarize(va), bond_10y_annualized=float(np.prod(1 + bret) ** (1 / 10) - 1))
    res["rate_scenarios"] = dict(y0=y0, table=scen)
    # 환 민감도
    fx = monthly_last("KRW_X").pct_change().dropna()
    fx.index = fx.index.to_period("M")
    K = P.join(fx.rename("fx"), how="inner")["2004-01":]
    krw = {}
    for name in ("A 주식100", "B 70/30(장기채)", "C 60/40(장기채)", "F 60/장기채30/금10"):
        tw_d = ALLOC[name]
        cols = list(tw_d)
        sub2 = K[cols + ["fx"]].dropna()
        Rk = np.column_stack([(1 + sub2[c].values) * (1 + sub2["fx"].values) - 1 for c in cols])
        tw = np.array([tw_d[c] for c in cols])
        krw[name] = dict(usd=summarize([path(sub2[cols].values, tw, s) for s in range(0, len(sub2) - MONTHS + 1)]),
                         krw=summarize([path(Rk, tw, s) for s in range(0, len(sub2) - MONTHS + 1)]))
    res["krw"] = krw
    res["last_paths"] = paths_last
    RES.write_text(json.dumps(res, ensure_ascii=False, indent=1, default=float), encoding="utf-8")
    write_excel(res)
    print("done", res["meta"])


def write_excel(res):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    hdr = PatternFill("solid", fgColor="1F3864")

    def head(ws, r, cols):
        for j, h in enumerate(cols, 1):
            c = ws.cell(row=r, column=j, value=h)
            c.fill, c.font = hdr, Font(color="FFFFFF", bold=True)
    ws = wb.active
    ws.title = "안내"
    for i, t in enumerate(["연금저축 월 50만원 x 10년(원금 6,000만원) 시뮬레이션. 자문 아님 — 과거 자료 기반 시나리오.",
                           "배분: A 주식100 · B 70/30 · C 60/40(30년채) · D 60/40(10년채) · E 60/40(현금) · F 60/30/10(금). 주식=SPY, 채권=정밀 프록시, USD 로컬 수익, 매수 편도 5bp, 납입금을 부족 자산에 배분(매도 없음).",
                           "최종액/원금 1.0 = 원금 그대로. 1.5 = 원금의 1.5배(수익 +50%). 원(KRW) 환산은 6,000만원 x 배수.",
                           "역사 창은 1993년 이후라 채권이 1990년대~2020년 금리 하락기의 혜택을 크게 받았다. 부트스트랩은 이를 섞지만 체제 구조는 보존하지 못한다. 금리 시나리오 시트가 미래 민감도다.",
                           "환율·세금·수수료는 반영하지 않았다(환 민감도는 별도 시트)."], 1):
        ws.cell(row=i, column=1, value=t)
    ws.column_dimensions["A"].width = 150
    ws = wb.create_sheet("1_역사롤링10년")
    ws["A1"] = f"역사 롤링 10년 창(매달 시작) — 최종액/원금 배수. 원(KRW)=배수 x {PRINCIPAL:,}원"
    ws["A1"].font = Font(bold=True)
    head(ws, 3, ["배분", "창 수", "최악", "하위10%", "중앙", "상위10%", "최고", "원금 미만 확률", "IRR 중앙", "IRR 최악", "기간중 최저(평가/납입)", "중앙(원)", "하위10%(원)", "최악(원)"])
    for i, (n, v) in enumerate(res["hist"].items(), 4):
        vals = [n, v["n"], v["worst"], v["p10"], v["median"], v["p90"], v["best"], v["loss_prob"], v["irr_median"], v["irr_worst"], v["underwater_min"],
                v["median"] * PRINCIPAL, v["p10"] * PRINCIPAL, v["worst"] * PRINCIPAL]
        for j, x in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=x)
            if j in (3, 4, 5, 6, 7, 11):
                c.number_format = "0.00"
            elif j in (8, 9, 10):
                c.number_format = "0.0%"
            elif j >= 12:
                c.number_format = "#,##0"
    bc = BarChart()
    bc.title = "최종액/원금 — 중앙 · 하위10% · 최악"
    for col in (5, 4, 3):
        bc.add_data(Reference(ws, min_col=col, min_row=3, max_row=3 + len(res["hist"])), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + len(res["hist"])))
    bc.height, bc.width = 9, 22
    ws.add_chart(bc, "A12")
    for j in range(1, 15):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws2 = wb.create_sheet("2_부트스트랩5000")
    ws2["A1"] = "블록 부트스트랩 5,000경로(12개월 블록) — 최종액/원금 배수"
    ws2["A1"].font = Font(bold=True)
    head(ws2, 3, ["배분", "최악", "하위10%", "중앙", "상위10%", "최고", "원금 미만 확률", "IRR 중앙", "IRR 하위10%", "중앙(원)", "하위10%(원)"])
    for i, (n, v) in enumerate(res["boot"].items(), 4):
        vals = [n, v["worst"], v["p10"], v["median"], v["p90"], v["best"], v["loss_prob"], v["irr_median"], v["irr_p10"], v["median"] * PRINCIPAL, v["p10"] * PRINCIPAL]
        for j, x in enumerate(vals, 1):
            c = ws2.cell(row=i, column=j, value=x)
            c.number_format = "0.00" if 2 <= j <= 6 else "0.0%" if j in (7, 8, 9) else "#,##0" if j >= 10 else "General"
    for j in range(1, 12):
        ws2.column_dimensions[get_column_letter(j)].width = 16
    ws3 = wb.create_sheet("3_시작시점별")
    ws3["A1"] = "시작 시점별 실제 사례 — 최종액(원) / 배수 / 기간 중 최저 평가배수"
    ws3["A1"].font = Font(bold=True)
    head(ws3, 3, ["배분"] + [f"{c}~ 최종(원)" for c in COHORTS] + [f"{c}~ 배수" for c in COHORTS] + [f"{c}~ 기간중 최저" for c in COHORTS])
    for i, (n, d) in enumerate(res["cohort"].items(), 4):
        ws3.cell(row=i, column=1, value=n)
        for k, cs in enumerate(COHORTS):
            v = d.get(cs)
            if v:
                ws3.cell(row=i, column=2 + k, value=v["final_krw"]).number_format = "#,##0"
                ws3.cell(row=i, column=2 + len(COHORTS) + k, value=v["final_ratio"]).number_format = "0.00"
                ws3.cell(row=i, column=2 + 2 * len(COHORTS) + k, value=v["underwater_min"]).number_format = "0.00"
    for j in range(1, 17):
        ws3.column_dimensions[get_column_letter(j)].width = 15
    ws4 = wb.create_sheet("4_금리시나리오")
    ws4["A1"] = f"30년물 금리 시나리오(현재 {res['rate_scenarios']['y0']:.2f}% 에서 10년에 걸쳐 직선 이동) — 배분 C(60/40 장기채) 최종액/원금, 주식은 실제 창 수익"
    ws4["A1"].font = Font(bold=True)
    head(ws4, 3, ["10년 뒤 30년물 금리%", "채권 다리 연수익", "C 최악", "C 하위10%", "C 중앙", "C 원금미만확률", "A(주식100) 중앙", "A 하위10%"])
    for i, (k, v) in enumerate(res["rate_scenarios"]["table"].items(), 4):
        vals = [float(k), v["bond_10y_annualized"], v["C"]["worst"], v["C"]["p10"], v["C"]["median"], v["C"]["loss_prob"], v["A"]["median"], v["A"]["p10"]]
        for j, x in enumerate(vals, 1):
            ws4.cell(row=i, column=j, value=x).number_format = "0.00" if j in (1, 3, 4, 5, 7, 8) else "0.0%"
    for j in range(1, 9):
        ws4.column_dimensions[get_column_letter(j)].width = 20
    ws5 = wb.create_sheet("5_환민감도")
    ws5["A1"] = "2004-01~ 창: 달러 기준(USD) 대 원화 비헤지 — 최종액/원금 중앙·하위10%·최악"
    ws5["A1"].font = Font(bold=True)
    head(ws5, 3, ["배분", "USD 중앙", "USD 하위10%", "USD 최악", "원화비헤지 중앙", "원화 하위10%", "원화 최악", "창 수"])
    for i, (n, d) in enumerate(res["krw"].items(), 4):
        vals = [n, d["usd"]["median"], d["usd"]["p10"], d["usd"]["worst"], d["krw"]["median"], d["krw"]["p10"], d["krw"]["worst"], d["usd"]["n"]]
        for j, x in enumerate(vals, 1):
            ws5.cell(row=i, column=j, value=x).number_format = "0.00" if 2 <= j <= 7 else "General"
    for j in range(1, 9):
        ws5.column_dimensions[get_column_letter(j)].width = 18
    ws6 = wb.create_sheet("6_실제지난10년")
    ws6["A1"] = "2016-10 시작~2026-09 (가장 최근 완결 10년) 월말 평가액(원). 원금은 월 50만원 누적"
    ws6["A1"].font = Font(bold=True)
    names = list(res["hist"])
    head(ws6, 3, ["개월", "누적 원금"] + names)
    for m in range(MONTHS):
        ws6.cell(row=4 + m, column=1, value=m + 1)
        ws6.cell(row=4 + m, column=2, value=MONTHLY * (m + 1)).number_format = "#,##0"
        for j, n in enumerate(names, 3):
            ws6.cell(row=4 + m, column=j, value=res["last_paths"][n + "_value"][m]).number_format = "#,##0"
    lc = LineChart()
    lc.title = "실제 지난 10년(2016-10~2026-09) 월 50만원 평가액"
    lc.add_data(Reference(ws6, min_col=2, max_col=2 + len(names), min_row=3, max_row=3 + MONTHS), titles_from_data=True)
    lc.set_categories(Reference(ws6, min_col=1, min_row=4, max_row=3 + MONTHS))
    lc.height, lc.width = 11, 24
    ws6.add_chart(lc, "K3")
    for j in range(1, 10):
        ws6.column_dimensions[get_column_letter(j)].width = 16
    wb.save(OUT / "pension-dca10y.xlsx")


def selftest():
    Rv = np.zeros((200, 2))
    V = path(Rv, np.array([.5, .5]), 0)
    assert abs(V[-1] - MONTHS * (1 - COST)) < 1e-6
    Rc = np.full((200, 1), 0.01)
    Vc = path(Rc, np.array([1.0]), 0)
    r = irr(Vc[-1])
    assert abs(r - ((1.01 * (1 - 0)) ** 12 - 1)) < 5e-3, r   # 비용 5bp 만큼 약간 낮음
    s = summarize([Vc, V])
    assert s["n"] == 2 and s["worst"] < s["best"]
    print("selftest ok")


if __name__ == "__main__":
    selftest() if "--selftest" in sys.argv else run()

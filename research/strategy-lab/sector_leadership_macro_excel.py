#!/usr/bin/env python3
"""국내·미국 섹터 주도 구간 + 환율·금리·S&P500·나스닥100 월별/연도별 엑셀 — 기술통계(판정 없음).

    python research/strategy-lab/sector_leadership_macro_excel.py --selftest
    python research/strategy-lab/sector_leadership_macro_excel.py
    -> research/strategy-lab/reports/2026-09-sector-leadership-macro/sector-leadership-macro.xlsx

정의(고정): 주도 = 월말 기준 직전 6개월(당월 포함) 누적수익 상위 섹터. 1위·2위·3위를 낸다. 이달 1위/꼴찌 = 그 달 수익 기준.
한국 = 20개 투자그룹(`config/sectorGroups.json` 롤업, 종목 등가중 월수익, 2016-02~2026-07 — 현재 업종분류·생존종목 근사, PIT 아님).
미국 = SPDR 9섹터(XLB XLE XLF XLI XLK XLP XLU XLV XLY, 1999-01~ 총수익, 통신·부동산 섹터 제외).
거시: 환율 FRED DEXKOUS(월말) · 한국 국고채 3년(FRED, ~2026-08) · 미국 3M/10Y/30Y(Yahoo ^IRX ^TNX ^TYX 월말) · 미국 기준금리(FRED) · KOSPI(FRED).
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd

LAB = Path(__file__).resolve().parent
sys.path.insert(0, str(LAB))
from pension_sector_bond_test import mret, kr_returns  # noqa: E402

MR = LAB / "data" / "market-regime"
OUT = LAB / "reports" / "2026-09-sector-leadership-macro"
US_NAMES = {"XLB": "소재", "XLE": "에너지", "XLF": "금융", "XLI": "산업재", "XLK": "기술", "XLP": "필수소비재",
            "XLU": "유틸리티", "XLV": "헬스케어", "XLY": "경기소비재"}


def raw_month_end(name, col="value", datecol="date"):
    d = pd.read_parquet(MR / f"{name}.parquet")
    s = pd.Series(d[col].values, index=pd.to_datetime(d[datecol]))
    return s.sort_index()


def month_series(s: pd.Series) -> pd.Series:
    s = s.sort_index()
    m = s.groupby(s.index.to_period("M")).last()
    return m


def macro_frame():
    fx = month_series(raw_month_end("dexkous_raw"))
    fxl = pd.read_parquet(LAB / "data" / "pension-test" / "KRW_X.parquet")["close"]
    fxl = month_series(fxl)
    fx = fx.combine_first(fxl[fxl.index > fx.index[-1]])
    kr3 = month_series(raw_month_end("krtreasury3y_raw"))
    ff = month_series(raw_month_end("usfedfundsrate_raw"))
    kospi = month_series(raw_month_end("krkospi_raw"))
    def yh(n):
        return month_series(pd.read_parquet(LAB / "data" / "pension-test" / f"{n}.parquet")["close"])
    df = pd.DataFrame({"usdkrw": fx, "kr3y": kr3, "ff": ff, "us3m": yh("_IRX"), "us10y": yh("_TNX"), "us30y": yh("_TYX"), "kospi": kospi})
    df.index = df.index.astype("period[M]")
    return df.sort_index()


def to_period(R: pd.DataFrame) -> pd.DataFrame:
    R = R.copy()
    R.index = pd.DatetimeIndex(R.index).to_period("M")
    return R


def leaders(R: pd.DataFrame, J=6):
    lg = np.log1p(R)
    tr = np.expm1(lg.rolling(J).sum())
    rank = tr.rank(axis=1, ascending=False, method="first")
    top = {k: rank.eq(k).idxmax(axis=1).where(rank.eq(k).any(axis=1)) for k in (1, 2, 3)}
    topv = {k: pd.Series([tr.loc[i, top[k][i]] if pd.notna(top[k][i]) else np.nan for i in tr.index], index=tr.index) for k in (1, 2, 3)}
    best = R.idxmax(axis=1)
    worst = R.idxmin(axis=1)
    return tr, top, topv, best, worst


def build_monthly(R: pd.DataFrame, macro: pd.DataFrame, us_bench: pd.DataFrame, names: dict, kr: bool):
    tr, top, topv, best, worst = leaders(R)
    idx = macro.index.union(R.index)
    idx = idx[(idx >= R.index[0])]
    out = []
    for p in idx:
        have = p in R.index
        g = lambda s: names.get(s, s) if isinstance(s, str) else None
        row = {"연월": str(p),
               "6개월 1위": g(top[1].get(p)) if have else None, "1위 6M수익": topv[1].get(p) if have else None,
               "6개월 2위": g(top[2].get(p)) if have else None, "2위 6M수익": topv[2].get(p) if have else None,
               "6개월 3위": g(top[3].get(p)) if have else None, "3위 6M수익": topv[3].get(p) if have else None,
               "이달 1위": g(best.get(p)) if have else None, "이달 1위 수익": R.loc[p].max() if have else None,
               "이달 꼴찌": g(worst.get(p)) if have else None, "이달 꼴찌 수익": R.loc[p].min() if have else None}
        if kr:
            k = macro["kospi"]
            row["KOSPI 월수익"] = (k.get(p) / k.get(p - 1) - 1) if (p in k.index and (p - 1) in k.index) else None
        sp = us_bench["SPY"].get(p) if p in us_bench.index else None
        qq = us_bench["QQQ"].get(p) if p in us_bench.index else None
        fxm = (macro["usdkrw"].get(p) / macro["usdkrw"].get(p - 1) - 1) if (p in macro.index and (p - 1) in macro.index and pd.notna(macro["usdkrw"].get(p)) and pd.notna(macro["usdkrw"].get(p - 1))) else None
        row.update({"S&P500 월수익(USD)": sp, "나스닥100 월수익(USD)": qq,
                    "S&P500(원화환산)": ((1 + sp) * (1 + fxm) - 1) if (sp is not None and pd.notna(sp) and fxm is not None) else None,
                    "나스닥100(원화환산)": ((1 + qq) * (1 + fxm) - 1) if (qq is not None and pd.notna(qq) and fxm is not None) else None,
                    "원/달러(월말)": macro["usdkrw"].get(p), "환율 월변화": fxm})
        for col, lab in (("kr3y", "한국 국고채3년%"), ("us3m", "미국 3개월%"), ("us10y", "미국 10년%"), ("us30y", "미국 30년%"), ("ff", "미국 기준금리%")):
            row[lab] = macro[col].get(p)
        for col, lab in (("kr3y", "국고3년 월변화(bp)"), ("us10y", "미국10년 월변화(bp)"), ("us30y", "미국30년 월변화(bp)")):
            a, b = macro[col].get(p), macro[col].get(p - 1)
            row[lab] = (a - b) * 100 if pd.notna(a) and pd.notna(b) else None
        out.append(row)
    return pd.DataFrame(out)


def build_runs(R: pd.DataFrame, macro: pd.DataFrame, us_bench: pd.DataFrame, names: dict, kr: bool):
    tr, top, topv, best, worst = leaders(R)
    L = top[1].dropna()
    runs, start = [], None
    prev, prev_p = None, None
    for p, s in L.items():
        if s != prev:
            if prev is not None:
                runs.append((prev, start, prev_p))
            prev, start = s, p
        prev_p = p
    runs.append((prev, start, prev_p))
    rows = []
    for s, a, b in runs:
        months = pd.period_range(a, b, freq="M")
        cum = float(np.prod(1 + R.loc[months, s]) - 1)
        def bench(x):
            v = us_bench[x].reindex(months)
            return float(np.prod(1 + v)) - 1 if v.notna().all() else None
        def chg(col, pct=False):
            x0, x1 = macro[col].get(a - 1), macro[col].get(b)
            if pd.isna(x0) or pd.isna(x1):
                return None
            return (x1 / x0 - 1) if pct else (x1 - x0) * 100
        avg_all = float(np.prod(1 + R.loc[months].mean(axis=1)) - 1)
        row = {"주도 섹터(6개월 1위)": names.get(s, s), "시작": str(a), "끝": str(b), "개월수": len(months),
               "구간 섹터수익": cum, "전체 섹터 평균": avg_all, "초과": cum - avg_all,
               "S&P500(USD)": bench("SPY"), "나스닥100(USD)": bench("QQQ"), "원/달러 변화": chg("usdkrw", True),
               "국고3년 변화(bp)": chg("kr3y"), "미국10년 변화(bp)": chg("us10y"), "미국30년 변화(bp)": chg("us30y"),
               "시작시 원/달러": macro["usdkrw"].get(a), "시작시 미국10년%": macro["us10y"].get(a)}
        if kr:
            k = macro["kospi"]
            row["KOSPI 구간수익"] = float(k.get(b) / k.get(a - 1) - 1) if (b in k.index and (a - 1) in k.index) else None
        rows.append(row)
    return pd.DataFrame(rows)


def build_yearly(R: pd.DataFrame, macro: pd.DataFrame, us_bench: pd.DataFrame, names: dict, kr: bool):
    tr, top, topv, best, worst = leaders(R)
    rows = []
    for y in sorted(set(p.year for p in R.index)):
        ps = [p for p in R.index if p.year == y]
        yr = (1 + R.loc[ps]).prod() - 1
        s = yr.sort_values(ascending=False)
        n1 = top[1].reindex(ps).dropna().value_counts()
        row = {"연도": y, "개월수": len(ps),
               "연 수익 1위": f"{names.get(s.index[0], s.index[0])} {s.iloc[0]:+.0%}", "2위": f"{names.get(s.index[1], s.index[1])} {s.iloc[1]:+.0%}",
               "3위": f"{names.get(s.index[2], s.index[2])} {s.iloc[2]:+.0%}",
               "꼴찌": f"{names.get(s.index[-1], s.index[-1])} {s.iloc[-1]:+.0%}",
               "섹터 평균": float((1 + R.loc[ps].mean(axis=1)).prod() - 1),
               "6개월1위 최다(개월)": (f"{names.get(n1.index[0], n1.index[0])} {n1.iloc[0]}개월" if len(n1) else None),
               "6개월1위 2번째": (f"{names.get(n1.index[1], n1.index[1])} {n1.iloc[1]}개월" if len(n1) > 1 else None)}
        def ann(v):
            x = v.reindex(ps)
            return float((1 + x).prod() - 1) if x.notna().all() else None
        row["S&P500(USD)"] = ann(us_bench["SPY"])
        row["나스닥100(USD)"] = ann(us_bench["QQQ"])
        e0, e1 = pd.Period(f"{y-1}-12"), pd.Period(f"{y}-12")
        e1 = ps[-1]
        fx0, fx1 = macro["usdkrw"].get(e0), macro["usdkrw"].get(e1)
        row["원/달러 연말"] = fx1
        row["원/달러 연변화"] = (fx1 / fx0 - 1) if pd.notna(fx0) and pd.notna(fx1) else None
        for col, lab in (("kr3y", "국고3년 연말%"), ("us3m", "미국3M 연말%"), ("us10y", "미국10Y 연말%"), ("us30y", "미국30Y 연말%"), ("ff", "미국기준금리 연말%")):
            row[lab] = macro[col].get(e1)
        for col, lab in (("kr3y", "국고3년 연변화(bp)"), ("us10y", "미국10Y 연변화(bp)")):
            a, b = macro[col].get(e0), macro[col].get(e1)
            row[lab] = (b - a) * 100 if pd.notna(a) and pd.notna(b) else None
        if kr:
            k = macro["kospi"]
            row["KOSPI 연수익"] = float(k.get(e1) / k.get(e0) - 1) if (e1 in k.index and e0 in k.index) else None
        rows.append(row)
    return pd.DataFrame(rows)


def build_regime(R: pd.DataFrame, macro: pd.DataFrame, names: dict):
    tr, top, topv, best, worst = leaders(R)
    L = top[1].dropna()
    rows = []
    for p, s in L.items():
        d10 = macro["us10y"].get(p) - macro["us10y"].get(p - 6) if pd.notna(macro["us10y"].get(p)) and pd.notna(macro["us10y"].get(p - 6)) else np.nan
        f = macro["usdkrw"].get(p) / macro["usdkrw"].get(p - 6) - 1 if pd.notna(macro["usdkrw"].get(p)) and pd.notna(macro["usdkrw"].get(p - 6)) else np.nan
        rows.append((p, s, d10, f))
    d = pd.DataFrame(rows, columns=["p", "s", "d10", "fx"]).dropna()
    d["금리(6개월)"] = np.where(d.d10 > 0.25, "미국10년 상승(+25bp 초과)", np.where(d.d10 < -0.25, "미국10년 하락(-25bp 미만)", "횡보"))
    d["환율(6개월)"] = np.where(d.fx > 0.03, "원화 약세(+3% 초과)", np.where(d.fx < -0.03, "원화 강세(-3% 미만)", "횡보"))
    out = []
    for lab, col in (("금리", "금리(6개월)"), ("환율", "환율(6개월)")):
        for reg, g in d.groupby(col):
            vc = g.s.value_counts()
            out.append({"구분": lab, "국면": reg, "월수": len(g),
                        "1위 빈도 1": f"{names.get(vc.index[0], vc.index[0])} {vc.iloc[0] / len(g):.0%}",
                        "2": (f"{names.get(vc.index[1], vc.index[1])} {vc.iloc[1] / len(g):.0%}" if len(vc) > 1 else None),
                        "3": (f"{names.get(vc.index[2], vc.index[2])} {vc.iloc[2] / len(g):.0%}" if len(vc) > 2 else None)})
    return pd.DataFrame(out)


def write_excel(path, sheets, notes):
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"안내": notes}).to_excel(xw, sheet_name="안내", index=False)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(path)
    ws = wb["안내"]
    ws.column_dimensions["A"].width = 150
    for r in ws.iter_rows(min_row=2):
        r[0].alignment = Alignment(wrap_text=True, vertical="top")
    hdr = PatternFill("solid", fgColor="1F3864")
    for name, df in sheets.items():
        w = wb[name]
        w.freeze_panes = "B2"
        for c in w[1]:
            c.fill, c.font, c.alignment = hdr, Font(color="FFFFFF", bold=True), Alignment(wrap_text=True, horizontal="center", vertical="center")
        w.row_dimensions[1].height = 34
        for j, col in enumerate(df.columns, 1):
            L = get_column_letter(j)
            w.column_dimensions[L].width = max(10, min(26, len(str(col)) * 1.6 + 2)) if df[col].dtype != object else 20
            pct = any(k in col for k in ("수익", "변화", "월변", "초과", "평균", "S&P", "나스닥", "KOSPI")) and "bp" not in col and "%" not in col.replace("수익", "")
            if "bp" in col:
                fmt = '0'
            elif "원/달러" in col and "변화" not in col:
                fmt = '#,##0.0'
            elif col.endswith("%") or "%)" in col:
                fmt = '0.00'
            elif pct:
                fmt = '0.0%'
            else:
                fmt = None
            if fmt:
                for cell in w[L][1:]:
                    cell.number_format = fmt
            if pct and len(df) > 2:
                w.conditional_formatting.add(f"{L}2:{L}{len(df) + 1}", ColorScaleRule(start_type="num", start_value=-0.15, start_color="F8696B",
                                             mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=0.15, end_color="63BE7B"))
    wb.save(path)


def main():
    macro = macro_frame()
    us_bench = pd.DataFrame({"SPY": mret("SPY"), "QQQ": mret("QQQ")})
    us_bench.index = us_bench.index.to_period("M")
    Rus = to_period(pd.concat({s: mret(s) for s in US_NAMES}, axis=1).dropna()["1999-01":])
    Rkr = to_period(kr_returns())
    sheets = {}
    for tag, R, names, kr in (("한국", Rkr, {}, True), ("미국", Rus, US_NAMES, False)):
        sheets[f"{tag}_월별"] = build_monthly(R, macro, us_bench, names, kr)
        sheets[f"{tag}_주도구간"] = build_runs(R, macro, us_bench, names, kr)
        sheets[f"{tag}_연도별"] = build_yearly(R, macro, us_bench, names, kr)
        sheets[f"{tag}_국면별"] = build_regime(R, macro, names)
        rr = R.copy()
        rr.index = rr.index.astype(str)
        rr.columns = [names.get(c, c) for c in rr.columns]
        sheets[f"{tag}_섹터월수익"] = rr.reset_index().rename(columns={"index": "연월", "date": "연월"})
    notes = [
        "주도 섹터 = 월말 기준 직전 6개월(당월 포함) 누적수익이 가장 높은 섹터. '이달 1위/꼴찌'는 그 달 수익 기준. 색: 수익 열은 -15%(빨강)~+15%(초록).",
        "한국: 20개 투자그룹(config/sectorGroups.json 롤업), 그룹 내 종목 등가중 월수익, 2016-02~2026-07(126개월). 현재 업종분류·생존 종목 기준이라 과거 PIT 분류가 아니며 폐지 종목이 빠져 있다. 2016 이전은 자료가 없다.",
        "미국: SPDR 섹터 9개(소재·에너지·금융·산업재·기술·필수소비재·유틸리티·헬스케어·경기소비재) 총수익, 1999-01~2026-09. 통신·부동산 섹터는 이력이 짧아 제외. 2018년 GICS 개편(통신 서비스 신설)으로 기술·경기소비재 구성이 바뀌었다.",
        "거시: 원/달러=FRED DEXKOUS 월말(끝난 뒤 달은 Yahoo KRW=X) · 한국 국고채3년=FRED(2026-08 까지) · 미국 3개월/10년/30년=Yahoo(^IRX ^TNX ^TYX) 월말 · 미국 기준금리=FRED · KOSPI=FRED(2014-05~). S&P500·나스닥100 = SPY·QQQ 총수익. 원화환산 = (1+지수수익)(1+환율변화)-1.",
        "주도구간 시트: 6개월 1위가 같은 섹터로 이어진 연속 구간(1개월짜리 포함). 구간 섹터수익 = 그 구간 개월 수익을 복리로 곱한 값, 전체 섹터 평균 대비 초과를 함께 냈다. 환율·금리 변화는 구간 시작 전월말→끝월말.",
        "국면별 시트: 6개월 1위 빈도를 '미국10년 6개월 변화(±25bp)'와 '원/달러 6개월 변화(±3%)' 국면별로 센 것. 참고용이며 인과가 아니다.",
        "★ 해석 주의: 앞선 시험(pension-sector-bond-results-2026-09)에서 주도 섹터는 6개월 뒤에도 주도일 확률이 우연과 같았고(지속성 근거 없음), 섹터 로테이션은 REJECT 였다. 이 표는 '과거에 무엇이 주도했나'의 기록이며 앞으로의 주도주를 알려주지 않는다.",
        "★ KOSPI 는 시가총액 가중이라 대형 반도체가 지수를 끌어올린다. 한국 섹터 그룹은 종목 등가중이라 지수와 크게 다를 수 있다(2025년 KOSPI +76% 대 섹터 평균 +13%). 두 수치를 같은 것으로 읽지 않는다.",
        "★ 한국 표본은 2025~26 국내 증시 급등(반도체·IT 중심)이 마지막 구간에 크게 들어 있다.",
    ]
    OUT.mkdir(parents=True, exist_ok=True)
    p = OUT / "sector-leadership-macro.xlsx"
    write_excel(p, sheets, notes)
    print("wrote", p, {k: v.shape for k, v in sheets.items()})


def selftest():
    idx = pd.period_range("2020-01", periods=12, freq="M")
    R = pd.DataFrame({"a": 0.05, "b": 0.01, "c": -0.01}, index=idx)
    tr, top, topv, best, worst = leaders(R)
    assert top[1].dropna().eq("a").all() and top[3].dropna().eq("c").all() and best.iloc[0] == "a"
    macro = pd.DataFrame({"usdkrw": 1000.0, "kr3y": 3.0, "ff": 1.0, "us3m": 1.0, "us10y": 2.0, "us30y": 3.0, "kospi": 2000.0},
                         index=pd.period_range("2019-06", periods=24, freq="M"))
    ub = pd.DataFrame({"SPY": 0.01, "QQQ": 0.02}, index=idx)
    runs = build_runs(R, macro, ub, {}, True)
    assert len(runs) == 1 and runs.loc[0, "개월수"] == 7
    assert len(build_yearly(R, macro, ub, {}, True)) == 1
    print("selftest ok")


if __name__ == "__main__":
    selftest() if "--selftest" in sys.argv else main()
